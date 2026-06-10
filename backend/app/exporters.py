from __future__ import annotations

import json
from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from myagent.backend.app.models import CompetitorReport


def report_to_markdown(report: CompetitorReport) -> str:
    if report.deep_dive_markdown.strip():
        return report.deep_dive_markdown.strip() + "\n"

    lines: list[str] = [
        f"# {report.target_product} 竞品分析报告",
        "",
        "## 执行摘要",
        report.executive_summary,
        "",
    ]
    if report.risk_notice:
        lines.extend(["## 风险提示", report.risk_notice, ""])

    lines.append("## 竞品画像")
    for profile in report.competitors:
        lines.extend(
            [
                f"### {profile.name}",
                f"- 定位：{profile.positioning}",
                f"- 目标用户：{', '.join(profile.target_users) or 'Unknown'}",
                f"- 核心功能：{', '.join(profile.core_features) or 'Unknown'}",
                f"- 定价：{profile.pricing}",
                f"- 商业模式：{profile.business_model}",
                f"- 渠道：{', '.join(profile.channels) or 'Unknown'}",
                f"- 分数：{profile.score if profile.score is not None else 'Unknown'}",
                f"- 选择理由：{profile.selection_reason or '未提供'}",
                f"- 证据：{', '.join(profile.evidence_ids) or 'None'}",
                "",
            ]
        )

    lines.append("## 关键结论")
    for finding in report.findings:
        lines.extend(
            [
                f"### {finding.title}",
                f"- 类别：{finding.category}",
                f"- 结论：{finding.conclusion}",
                f"- 置信度：{finding.confidence}",
                f"- 影响：{finding.impact}",
                f"- 证据：{', '.join(finding.evidence_ids) or 'None'}",
                "",
            ]
        )

    lines.extend(
        [
            "## SWOT",
            f"- 优势：{'; '.join(report.swot.strengths) or 'None'}",
            f"- 劣势：{'; '.join(report.swot.weaknesses) or 'None'}",
            f"- 机会：{'; '.join(report.swot.opportunities) or 'None'}",
            f"- 威胁：{'; '.join(report.swot.threats) or 'None'}",
            "",
            "## 建议",
        ]
    )
    lines.extend(f"- {item}" for item in report.recommendations)
    lines.extend(["", "## 来源"])
    for item in report.evidence:
        image = f"；截图：{item.image_path}" if item.image_path else ""
        lines.extend([f"- [{item.id}] {item.title}: {item.url}{image}", f"  - 摘录：{item.excerpt[:240]}"])
    return "\n".join(lines).strip() + "\n"


def report_to_json(report: CompetitorReport) -> bytes:
    return json.dumps(report.model_dump(mode="json"), ensure_ascii=False, indent=2).encode("utf-8")


def report_to_pdf(report: CompetitorReport) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, title=f"{report.target_product} 竞品分析报告")
    styles = getSampleStyleSheet()
    for style_name in styles.byName:
        styles[style_name].fontName = "STSong-Light"

    story = []
    markdown = report_to_markdown(report)
    for line in markdown.splitlines():
        if line.startswith("# "):
            story.append(Paragraph(line[2:], styles["Title"]))
        elif line.startswith("## "):
            story.append(Spacer(1, 10))
            story.append(Paragraph(line[3:], styles["Heading2"]))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], styles["Heading3"]))
        elif line.strip():
            story.append(Paragraph(line.replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"]))
        else:
            story.append(Spacer(1, 6))
    doc.build(story)
    return buffer.getvalue()


def report_to_docx(report: CompetitorReport) -> bytes:
    return markdown_to_docx(report_to_markdown(report), report.target_product)


def markdown_to_docx(markdown: str, title: str) -> bytes:
    document = Document()
    document.add_heading(f"{title} 竞品分析报告", level=1)
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("#### "):
            document.add_heading(line[5:], level=4)
        elif line.startswith("- [ ] "):
            document.add_paragraph(line[6:], style="List Bullet")
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_report(report: CompetitorReport, fmt: str) -> tuple[bytes, str, str]:
    if fmt == "md":
        return report_to_markdown(report).encode("utf-8"), "text/markdown; charset=utf-8", "md"
    if fmt == "json":
        return report_to_json(report), "application/json", "json"
    if fmt == "pdf":
        return report_to_pdf(report), "application/pdf", "pdf"
    if fmt == "docx":
        return (
            report_to_docx(report),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        )
    raise ValueError(f"Unsupported export format: {fmt}")
